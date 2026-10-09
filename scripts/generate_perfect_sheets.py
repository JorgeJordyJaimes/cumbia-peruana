import zipfile
import shutil
import xml.etree.ElementTree as ET
import re
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

ODS_PATH = 'docs/datos-investigacion/Datos BD.ods'
XLSX_PATH = 'docs/datos-investigacion/Datos BD.xlsx'
BACKUP_PATH = 'docs/datos-investigacion/Datos BD.ods.bak'

shutil.copyfile(ODS_PATH, BACKUP_PATH)
print(f'Backup de seguridad creado en: {BACKUP_PATH}')

# ==============================================================================
# 1. EXTRAER DATOS MAESTROS DE LOS ARCHIVOS SQL
# ==============================================================================
def extract_sellos():
    sellos = []
    with open('supabase/data/base/sellos.sql', 'r', encoding='utf-8') as f:
        content = f.read()
    for m in re.finditer(r"\('([^']+)'\)", content):
        sellos.append(m.group(1).strip())
    return sorted(list(dict.fromkeys(sellos)))

def extract_personas():
    personas = []
    with open('supabase/data/base/personas.sql', 'r', encoding='utf-8') as f:
        content = f.read()
    for m in re.finditer(r"\('([^']+)',\s*'([^']*)'\)", content):
        personas.append(m.group(1).strip())
    return sorted(list(dict.fromkeys(personas)))

def extract_grupos():
    grupos = []
    with open('supabase/data/base/grupos.sql', 'r', encoding='utf-8') as f:
        content = f.read()
    for m in re.finditer(r"\('((?:''|[^'])+)'(?:,\s*\d+)?\)", content):
        grupos.append(m.group(1).replace("''", "'").strip())
    return sorted(list(dict.fromkeys(grupos)))

sellos_list = extract_sellos()
personas_list = extract_personas()
grupos_list = extract_grupos()

camelot_list = [
    '1A', '1B', '2A', '2B', '3A', '3B', '4A', '4B',
    '5A', '5B', '6A', '6B', '7A', '7B', '8A', '8B',
    '9A', '9B', '10A', '10B', '11A', '11B', '12A', '12B'
]
sino_list = ['SI', 'NO']
lados_list = ['A', 'B']
lados_lp_list = ['Ambos', 'Lado A', 'Lado B']
generos_list = [
    'Cumbia Costeña', 'Cumbia Andina', 'Cumbia Amazónica', 'Chicha',
    'Cumbia Norteña', 'Cumbia Sureña', 'Cumbia Psicodélica', 'Tecnocumbia'
]

print(f'Maestros extraídos -> Sellos: {len(sellos_list)} | Personas: {len(personas_list)} | Grupos: {len(grupos_list)}')

# ==============================================================================
# 2. DEFINICIÓN ESTRUCTURAL DE LAS 10 HOJAS DE TRABAJO
# ==============================================================================

SHEETS_DEF = {
    'Ingreso Singles': {
        'type': 'ingreso',
        'headers': [
            'GRUPO / ARTISTA', 'SELLO DISCOGRÁFICO', 'NÚMERO DE CATÁLOGO', 'AÑO DE PUBLICACIÓN',
            '¿ES DISCO SPLIT?', '¿SOLO EN 45?', '¿INCLUIDO EN LP?', 'LADOS EN LP',
            '¿EXTRAÍDO DE LP?', 'LP RELACIONADO (CATÁLOGO)', '¿ES REEDICIÓN?', 'COMENTARIO / NOTAS'
        ],
        'widths': [
            '5.5cm', '4.5cm', '4.0cm', '3.5cm',
            '3.5cm', '3.2cm', '3.5cm', '3.2cm',
            '3.5cm', '4.5cm', '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Grupo o titular: menú desplegable)',
            '(Disquera: menú desplegable de 258 sellos)',
            '(Catálogo con ceros adelante: ej. 009, 0009. Texto protegido)',
            '(Año de salida: 1968 a 2005. Ej: 1973)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Ambos, Lado A, Lado B - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Catálogo del LP matriz si aplica)',
            '(SI / NO - Menú desplegable)',
            '(Notas sobre vinilo, dedicatorias o rarezas)'
        ],
        'examples': [
            ['Grupo Celeste', 'Difa', '009', '1973', 'NO', 'SI', 'NO', '', 'NO', '', 'NO', 'Single exclusivo (catálogo 009 protegido)'],
            ['Los Shapis', 'Horóscopo', '0045', '1981', 'NO', 'NO', 'SI', 'Ambos', 'NO', 'HLP-1001', 'NO', 'Catálogo con ceros preservados (0045)'],
            ['Los Ecos', 'Sonoradio', '13250', '1972', 'SI', 'SI', 'NO', '', 'NO', '', 'NO', 'Split: Lado A Los Ecos, Lado B Diablos Rojos']
        ],
        'validations': {
            0: 'val_grupos',
            1: 'val_sellos',
            4: 'val_sino',
            5: 'val_sino',
            6: 'val_sino',
            7: 'val_lados_lp',
            8: 'val_sino',
            10: 'val_sino'
        }
    },
    'Ingreso LP': {
        'type': 'ingreso',
        'headers': [
            'GRUPO / ARTISTA', 'SELLO DISCOGRÁFICO', 'NOMBRE DEL ÁLBUM', 'NÚMERO DE CATÁLOGO',
            'AÑO DE PUBLICACIÓN', '¿ES RECOPILATORIO?', '¿ES VARIOS ARTISTAS?', '¿ES DISCO SPLIT?',
            '¿ES REEDICIÓN?', 'ÁLBUM ORIGINAL (CATÁLOGO)', 'COMENTARIO / NOTAS'
        ],
        'widths': [
            '5.5cm', '4.5cm', '6.5cm', '4.0cm',
            '3.5cm', '3.8cm', '3.8cm', '3.5cm',
            '3.2cm', '4.5cm', '11.0cm'
        ],
        'guide': [
            '(Grupo principal o titular: menú desplegable)',
            '(Disquera: menú desplegable de 258 sellos)',
            '(Título completo del LP de larga duración. Ej: Constelación)',
            '(Catálogo con ceros si aplica. Ej: LPN-2415, HLP-0010)',
            '(Año de salida: 1968 a 2005. Ej: 1971)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Catálogo de la 1ª edición si es reedición)',
            '(Funda troquelada, tiraje, notas de matriz)'
        ],
        'examples': [
            ['Los Destellos', 'Odeón', 'Constelación', 'LPN-2415', '1971', 'NO', 'NO', 'NO', 'NO', '', '1ra edición con funda troquelada'],
            ['Varios Artistas', 'Horóscopo', 'Super Cumbiones Vol. 1', 'HLP-1015', '1983', 'SI', 'SI', 'NO', 'NO', '', 'Compilatorio con Chacalón, Shapis, Maravilla'],
            ['Los Destellos', 'Iempsa', 'Mano a Mano Tropical', 'ELD-02.01', '1974', 'NO', 'NO', 'SI', 'NO', '', 'Lado A: Los Destellos / Lado B: Manzanita']
        ],
        'validations': {
            0: 'val_grupos',
            1: 'val_sellos',
            5: 'val_sino',
            6: 'val_sino',
            7: 'val_sino',
            8: 'val_sino'
        }
    },
    'Ingreso EP': {
        'type': 'ingreso',
        'headers': [
            'GRUPO / ARTISTA', 'SELLO DISCOGRÁFICO', 'NOMBRE DEL EP', 'NÚMERO DE CATÁLOGO',
            'AÑO DE PUBLICACIÓN', '¿ES DISCO SPLIT?', '¿ES RECOPILATORIO?', '¿ES REEDICIÓN?',
            'COMENTARIO / NOTAS'
        ],
        'widths': [
            '5.5cm', '4.5cm', '6.5cm', '4.0cm',
            '3.5cm', '3.5cm', '3.8cm', '3.2cm',
            '11.0cm'
        ],
        'guide': [
            '(Grupo o titular: menú desplegable)',
            '(Disquera: menú desplegable de 258 sellos)',
            '(Título del EP de 4 temas. Ej: El Milagro Verde)',
            '(Catálogo con ceros si aplica. Ej: EP-001, EP-015)',
            '(Año de publicación. Ej: 1974)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Velocidad 33 o 45 RPM, notas de empaque)'
        ],
        'examples': [
            ['Los Mirlos', 'Infopesa', 'El Milagro Verde', 'EP-015', '1974', 'NO', 'NO', 'NO', 'Extended Play de 4 temas a 33 RPM']
        ],
        'validations': {
            0: 'val_grupos',
            1: 'val_sellos',
            5: 'val_sino',
            6: 'val_sino',
            7: 'val_sino'
        }
    },
    'Ingreso Cassete': {
        'type': 'ingreso',
        'headers': [
            'GRUPO / ARTISTA', 'SELLO DISCOGRÁFICO', 'NOMBRE DEL ÁLBUM', 'NÚMERO DE CATÁLOGO',
            'AÑO DE PUBLICACIÓN', '¿ES RECOPILATORIO?', '¿ES VARIOS ARTISTAS?', '¿ES CASSETTE SPLIT?',
            '¿ES REEDICIÓN?', 'COMENTARIO / NOTAS'
        ],
        'widths': [
            '5.5cm', '4.5cm', '6.5cm', '4.0cm',
            '3.5cm', '3.8cm', '3.8cm', '3.5cm',
            '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Grupo o titular: menú desplegable)',
            '(Disquera: menú desplegable de 258 sellos)',
            '(Título del álbum en casete. Ej: Llora... Llora con Los Destellos)',
            '(Catálogo con ceros si aplica. Ej: HK-0077, HK-1077)',
            '(Año de salida comercial en cinta. Ej: 1988)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Insert plegable, color de la cinta, rarezas)'
        ],
        'examples': [
            ['Los Destellos', 'Horóscopo', 'Llora... Llora con Los Destellos', 'HK-1077', '1988', 'NO', 'NO', 'NO', 'NO', 'Casete original con insert desplegable'],
            ['La Nueva Crema', 'Horóscopo', 'Éxitos de Oro en Casete', 'HK-1050', '1986', 'SI', 'NO', 'NO', 'NO', 'Recopilación oficial exclusiva en cinta']
        ],
        'validations': {
            0: 'val_grupos',
            1: 'val_sellos',
            5: 'val_sino',
            6: 'val_sino',
            7: 'val_sino',
            8: 'val_sino'
        }
    },
    'Ingreso CD': {
        'type': 'ingreso',
        'headers': [
            'GRUPO / ARTISTA', 'SELLO DISCOGRÁFICO', 'NOMBRE DEL ÁLBUM', 'NÚMERO DE CATÁLOGO',
            'AÑO DE PUBLICACIÓN', '¿ES RECOPILATORIO?', '¿ES VARIOS ARTISTAS?', '¿ES REEDICIÓN DE LP?',
            'LP ORIGINAL (CATÁLOGO)', 'COMENTARIO / NOTAS'
        ],
        'widths': [
            '5.5cm', '4.5cm', '6.5cm', '4.0cm',
            '3.5cm', '3.8cm', '3.8cm', '3.5cm',
            '4.5cm', '11.0cm'
        ],
        'guide': [
            '(Grupo o titular: menú desplegable)',
            '(Disquera: menú desplegable de 258 sellos)',
            '(Título del disco compacto. Ej: A Todo Ritmo)',
            '(Catálogo con ceros si aplica. Ej: CD-0012, CD-9012)',
            '(Año de publicación en CD: 1990 a 2005. Ej: 1995)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Catálogo del LP matriz original si fue remasterizado)',
            '(Caja jewel case, digipak, libreto)'
        ],
        'examples': [
            ['Armonia 10', 'Infopesa', 'A Todo Ritmo', 'CD-9012', '1995', 'NO', 'NO', 'NO', '', 'Edición en caja jewel case'],
            ['Los Destellos', 'Iempsa', '30 Años de Éxitos', 'CD-0105', '1998', 'SI', 'NO', 'SI', 'LPN-2415', 'Remasterización digital de cintas de 1971']
        ],
        'validations': {
            0: 'val_grupos',
            1: 'val_sellos',
            5: 'val_sino',
            6: 'val_sino',
            7: 'val_sino'
        }
    },
    'Temas Singles': {
        'type': 'temas',
        'headers': [
            'CATÁLOGO DEL SINGLE', 'LADO FÍSICO', 'PISTA N°', 'TÍTULO DE LA CANCIÓN',
            'COMPOSITOR REAL', 'CRÉDITO EN GALLETA (SEUDÓNIMO)', 'GRUPO INTÉRPRETE',
            'GÉNERO MUSICAL', 'DURACIÓN (MM:SS)', 'BPM (TEMPO)', 'TONO / CAMELOT',
            '¿ES MOSAICO?', 'LETRA / NOTAS DEL TEMA'
        ],
        'widths': [
            '4.2cm', '3.0cm', '2.8cm', '6.5cm',
            '5.5cm', '5.5cm', '5.5cm',
            '4.5cm', '3.5cm', '3.0cm', '3.5cm',
            '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Catálogo con ceros si aplica: ej. 009, 0009. Texto protegido)',
            '(A / B - Menú desplegable)',
            '(En 45 RPM siempre es 1 por cada lado)',
            '(Nombre del tema. Ej: En el Campo)',
            '(Compositor: menú desplegable de 360 personas)',
            '(Nombre o alias acreditado en el vinilo si difiere del real)',
            '(Grupo intérprete: menú desplegable de 433 grupos)',
            '(Género: menú desplegable de vertientes de cumbia)',
            '(Duración en minutos:segundos ej. 02:50 - TEXTO LIBRE)',
            '(Tempo en pulsos por minuto. Ej: 125 - NÚMERO LIBRE)',
            '(Camelot: menú desplegable 1A..12B)',
            '(SI / NO - Menú desplegable)',
            '(Voz solista, arreglos o notas de grabación - TEXTO LIBRE)'
        ],
        'examples': [
            ['009', 'A', '1', 'En el Campo', 'Víctor Homero Casahuamán Marín', 'V. Casahuamán', 'Grupo Celeste', 'Cumbia Costeña', '02:50', '125', '7A', 'NO', 'Voz solista: Lorenzo Palacios (Chacalón)'],
            ['009', 'B', '1', 'Melodía Celeste', 'Víctor Homero Casahuamán Marín', 'V. Casahuamán', 'Grupo Celeste', 'Cumbia Costeña', '02:45', '126', '7A', 'NO', 'Instrumental con guitarra eléctrica'],
            ['0045', 'A', '1', 'Parranda Chicha N° 1', 'Jaime Venturo Moreira Mercado', 'J. Moreira', 'Los Shapis', 'Chicha', '03:10', '102', '9A', 'SI', 'Parranda de dos temas en Lado A']
        ],
        'validations': {
            1: 'val_lados',
            4: 'val_personas',
            6: 'val_grupos',
            7: 'val_generos',
            10: 'val_camelot',
            11: 'val_sino'
        }
    },
    'Temas LP': {
        'type': 'temas',
        'headers': [
            'CATÁLOGO DEL LP', 'LADO FÍSICO', 'PISTA N°', 'TÍTULO DE LA CANCIÓN',
            'COMPOSITOR REAL', 'CRÉDITO EN GALLETA (SEUDÓNIMO)', 'GRUPO INTÉRPRETE',
            'GÉNERO MUSICAL', 'DURACIÓN (MM:SS)', 'BPM (TEMPO)', 'TONO / CAMELOT',
            '¿ES MOSAICO?', '¿ES GRABACIÓN INÉDITA?', 'SINGLE ORIGEN (CATÁLOGO)',
            'LETRA / NOTAS DEL TEMA'
        ],
        'widths': [
            '4.2cm', '3.0cm', '2.8cm', '6.5cm',
            '5.5cm', '5.5cm', '5.5cm',
            '4.5cm', '3.5cm', '3.0cm', '3.5cm',
            '3.2cm', '3.5cm', '4.2cm',
            '11.0cm'
        ],
        'guide': [
            '(Catálogo del LP registrado en Ingreso LP. Ej: HLP-1001)',
            '(A / B - Menú desplegable)',
            '(Número de surco: 1, 2, 3, 4, 5, 6...)',
            '(Nombre del tema. Ej: El Aguajal, Elsa)',
            '(Compositor: menú desplegable de 360 personas)',
            '(Nombre o alias acreditado en el vinilo si difiere del real)',
            '(Grupo intérprete: menú desplegable de 433 grupos)',
            '(Género: menú desplegable de vertientes de cumbia)',
            '(Duración en minutos:segundos ej. 03:15 - TEXTO LIBRE)',
            '(Tempo en pulsos por minuto. Ej: 98, 128 - NÚMERO LIBRE)',
            '(Camelot: menú desplegable 1A..12B)',
            '(SI / NO - Menú desplegable)',
            '(SI / NO - Menú desplegable)',
            '(Catálogo del 45 RPM si el tema proviene de un single previo)',
            '(Fragmento lírico destacado, solista o notas - TEXTO LIBRE)'
        ],
        'examples': [
            ['HLP-1001', 'A', '1', 'El Aguajal', 'Trinidad Hermoza', 'T. Hermoza', 'Los Shapis', 'Cumbia Andina', '03:15', '98', '9A', 'NO', 'NO', '0045', 'Voz: Julio Simeón (Chapulín el Dulce)'],
            ['HLP-1001', 'A', '2', 'Borrachito Borrachón', 'Jaime Venturo Moreira Mercado', 'J. Moreira', 'Los Shapis', 'Cumbia Andina', '03:22', '100', '9A', 'NO', 'NO', '', 'Guitarra con pedal de distorsión fuzz'],
            ['HLP-1001', 'B', '1', 'El Chofercito', 'Jaime Venturo Moreira Mercado', 'J. Moreira', 'Los Shapis', 'Cumbia Andina', '03:40', '104', '8A', 'NO', 'NO', '0045', 'Tema dedicado a transportistas'],
            ['HLP-1015', 'A', '4', 'Mosaico Tropical', 'Enrique Delgado Montes', 'D.R.', 'Los Destellos', 'Cumbia Costeña', '05:30', '128', '8A', 'SI', 'NO', '', 'Enganchado continuo de 3 temas clásicos']
        ],
        'validations': {
            1: 'val_lados',
            4: 'val_personas',
            6: 'val_grupos',
            7: 'val_generos',
            10: 'val_camelot',
            11: 'val_sino',
            12: 'val_sino'
        }
    },
    'Temas EP': {
        'type': 'temas',
        'headers': [
            'CATÁLOGO DEL EP', 'LADO FÍSICO', 'PISTA N°', 'TÍTULO DE LA CANCIÓN',
            'COMPOSITOR REAL', 'CRÉDITO EN GALLETA (SEUDÓNIMO)', 'GRUPO INTÉRPRETE',
            'GÉNERO MUSICAL', 'DURACIÓN (MM:SS)', 'BPM (TEMPO)', 'TONO / CAMELOT',
            '¿ES MOSAICO?', 'LETRA / NOTAS DEL TEMA'
        ],
        'widths': [
            '4.2cm', '3.0cm', '2.8cm', '6.5cm',
            '5.5cm', '5.5cm', '5.5cm',
            '4.5cm', '3.5cm', '3.0cm', '3.5cm',
            '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Catálogo del EP registrado en Ingreso EP. Ej: EP-015)',
            '(A / B - Menú desplegable)',
            '(1 o 2 por cada cara del Extended Play)',
            '(Nombre del tema. Ej: El Milagro Verde)',
            '(Compositor: menú desplegable de 360 personas)',
            '(Nombre o alias acreditado en la galleta)',
            '(Grupo intérprete: menú desplegable de 433 grupos)',
            '(Género: menú desplegable de vertientes de cumbia)',
            '(Duración en minutos:segundos - TEXTO LIBRE)',
            '(Tempo en pulsos por minuto - NÚMERO LIBRE)',
            '(Camelot: menú desplegable 1A..12B)',
            '(SI / NO - Menú desplegable)',
            '(Notas de instrumentación, percusión - TEXTO LIBRE)'
        ],
        'examples': [
            ['EP-015', 'A', '1', 'El Milagro Verde', 'Jorge Rodríguez Grández', 'J. Rodríguez', 'Los Mirlos', 'Cumbia Amazónica', '02:45', '130', '8A', 'NO', 'Guitarra amazónica con eco y wah-wah'],
            ['EP-015', 'A', '2', 'Lamento en la Selva', 'Jorge Rodríguez Grández', 'J. Rodríguez', 'Los Mirlos', 'Cumbia Amazónica', '02:50', '128', '8A', 'NO', 'Instrumental selvático']
        ],
        'validations': {
            1: 'val_lados',
            4: 'val_personas',
            6: 'val_grupos',
            7: 'val_generos',
            10: 'val_camelot',
            11: 'val_sino'
        }
    },
    'Temas Cassete': {
        'type': 'temas',
        'headers': [
            'CATÁLOGO DEL CASETE', 'LADO FÍSICO', 'PISTA N°', 'TÍTULO DE LA CANCIÓN',
            'COMPOSITOR REAL', 'CRÉDITO EN INSERT / CINTA', 'GRUPO INTÉRPRETE',
            'GÉNERO MUSICAL', 'DURACIÓN (MM:SS)', 'BPM (TEMPO)', 'TONO / CAMELOT',
            '¿ES MOSAICO?', 'LETRA / NOTAS DEL TEMA'
        ],
        'widths': [
            '4.2cm', '3.0cm', '2.8cm', '6.5cm',
            '5.5cm', '5.5cm', '5.5cm',
            '4.5cm', '3.5cm', '3.0cm', '3.5cm',
            '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Catálogo del casete registrado en Ingreso Cassete. Ej: HK-1077)',
            '(A / B - Menú desplegable)',
            '(Número de tema en ese lado: 1, 2, 3, 4, 5, 6...)',
            '(Nombre del tema grabado)',
            '(Compositor: menú desplegable de 360 personas)',
            '(Nombre o seudónimo en la cajita o cinta)',
            '(Grupo intérprete: menú desplegable de 433 grupos)',
            '(Género: menú desplegable de vertientes de cumbia)',
            '(Duración en minutos:segundos - TEXTO LIBRE)',
            '(Tempo en pulsos por minuto - NÚMERO LIBRE)',
            '(Camelot: menú desplegable 1A..12B)',
            '(SI / NO - Menú desplegable)',
            '(Notas líricas o solista vocal - TEXTO LIBRE)'
        ],
        'examples': [
            ['HK-1077', 'A', '1', 'Llora Corazón', 'Enrique Delgado Montes', 'E. Delgado', 'Los Destellos', 'Cumbia Costeña', '03:10', '126', '7A', 'NO', 'Grabación original en cinta'],
            ['HK-1077', 'A', '2', 'Mosaico Chicha en Vivo', 'Lorenzo Palacios Quispe', 'D.R.', 'Los Destellos', 'Chicha', '06:15', '115', '9A', 'SI', 'Potpourri continuo de éxitos en vivo']
        ],
        'validations': {
            1: 'val_lados',
            4: 'val_personas',
            6: 'val_grupos',
            7: 'val_generos',
            10: 'val_camelot',
            11: 'val_sino'
        }
    },
    'Temas CD': {
        'type': 'temas',
        'headers': [
            'CATÁLOGO DEL CD', 'PISTA N°', 'TÍTULO DE LA CANCIÓN',
            'COMPOSITOR REAL', 'CRÉDITO EN LIBRILLO (SEUDÓNIMO)', 'GRUPO INTÉRPRETE',
            'GÉNERO MUSICAL', 'DURACIÓN (MM:SS)', 'BPM (TEMPO)', 'TONO / CAMELOT',
            '¿ES MOSAICO?', 'LETRA / NOTAS DEL TEMA'
        ],
        'widths': [
            '4.2cm', '2.8cm', '6.5cm',
            '5.5cm', '5.5cm', '5.5cm',
            '4.5cm', '3.5cm', '3.0cm', '3.5cm',
            '3.2cm', '11.0cm'
        ],
        'guide': [
            '(Catálogo del CD registrado en Ingreso CD. Ej: CD-9012)',
            '(Pista numérica correlativa: 1, 2, 3... sin Lados)',
            '(Nombre del tema grabado)',
            '(Compositor: menú desplegable de 360 personas)',
            '(Nombre o crédito en el folleto/inlay del CD)',
            '(Grupo intérprete: menú desplegable de 433 grupos)',
            '(Género: menú desplegable de vertientes de cumbia)',
            '(Duración en minutos:segundos - TEXTO LIBRE)',
            '(Tempo en pulsos por minuto - NÚMERO LIBRE)',
            '(Camelot: menú desplegable 1A..12B)',
            '(SI / NO - Menú desplegable)',
            '(Notas sobre remasterización digital - TEXTO LIBRE)'
        ],
        'examples': [
            ['CD-9012', '1', 'El Cervecero', 'Enrique Delgado Montes', 'E. Delgado', 'Armonia 10', 'Cumbia Norteña', '03:45', '105', '8A', 'NO', 'Vocalista: Makuko Gallardo'],
            ['CD-9012', '2', 'Pagarás', 'Enrique Delgado Montes', 'E. Delgado', 'Armonia 10', 'Cumbia Norteña', '03:30', '108', '9A', 'NO', 'Sección de vientos y timbales caribeños']
        ],
        'validations': {
            3: 'val_personas',
            5: 'val_grupos',
            6: 'val_generos',
            9: 'val_camelot',
            10: 'val_sino'
        }
    }
}

# ==============================================================================
# 3. CONSTRUCCIÓN DE DATOS BD.ODS (LIBREOFFICE CALC COMPATIBLE)
# ==============================================================================
NS = {
    'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
    'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
    'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0',
    'style': 'urn:oasis:names:tc:opendocument:xmlns:style:1.0',
    'fo': 'urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0',
    'config': 'urn:oasis:names:tc:opendocument:xmlns:config:1.0',
    'calcext': 'urn:org:documentfoundation:names:experimental:calc:xmlns:calcext:1.0',
    'number': 'urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0',
    'of': 'urn:oasis:names:tc:opendocument:xmlns:of:1.2',
}

for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)

with zipfile.ZipFile(ODS_PATH, 'r') as z:
    content_xml = z.read('content.xml')
    settings_xml = z.read('settings.xml') if 'settings.xml' in z.namelist() else None
    other_files = {name: z.read(name) for name in z.namelist() if name not in ['content.xml', 'settings.xml']}

root = ET.fromstring(content_xml)

auto_styles = root.find('.//office:automatic-styles', NS)
if auto_styles is None:
    auto_styles = ET.SubElement(root, f"{{{NS['office']}}}automatic-styles")

if auto_styles.find(f"{{{NS['number']}}}text-style[@{{{NS['style']}}}name='NText']", NS) is None:
    t_style = ET.SubElement(auto_styles, f"{{{NS['number']}}}text-style", {
        f"{{{NS['style']}}}name": 'NText'
    })
    ET.SubElement(t_style, f"{{{NS['number']}}}text-content")

def create_cell_style(name, bg_color, text_color, bold=False, italic=False, font_size='9.5pt', is_text_format=False):
    s = auto_styles.find(f"{{{NS['style']}}}style[@{{{NS['style']}}}name='{name}']", NS)
    if s is None:
        s_attribs = {
            f"{{{NS['style']}}}name": name,
            f"{{{NS['style']}}}family": 'table-cell',
            f"{{{NS['style']}}}parent-style-name": 'Default'
        }
        if is_text_format:
            s_attribs[f"{{{NS['style']}}}data-style-name"] = 'NText'
        s = ET.SubElement(auto_styles, f"{{{NS['style']}}}style", s_attribs)
    else:
        if is_text_format:
            s.attrib[f"{{{NS['style']}}}data-style-name"] = 'NText'
            
    props = s.find(f"{{{NS['style']}}}table-cell-properties", NS)
    if props is None:
        props = ET.SubElement(s, f"{{{NS['style']}}}table-cell-properties")
    props.attrib[f"{{{NS['fo']}}}background-color"] = bg_color
    props.attrib[f"{{{NS['fo']}}}padding"] = '0.15cm'
    props.attrib[f"{{{NS['fo']}}}wrap-option"] = 'wrap'

    t_props = s.find(f"{{{NS['style']}}}text-properties", NS)
    if t_props is None:
        t_props = ET.SubElement(s, f"{{{NS['style']}}}text-properties")
    t_props.attrib[f"{{{NS['fo']}}}color"] = text_color
    t_props.attrib[f"{{{NS['style']}}}font-name"] = 'Arial'
    t_props.attrib[f"{{{NS['fo']}}}font-size"] = font_size
    if bold:
        t_props.attrib[f"{{{NS['fo']}}}font-weight"] = 'bold'
    if italic:
        t_props.attrib[f"{{{NS['fo']}}}font-style"] = 'italic'

create_cell_style('HeaderCellA', '#0f172a', '#ffffff', bold=True, font_size='10pt') # Navy
create_cell_style('HeaderCellB', '#1e3a8a', '#ffffff', bold=True, font_size='10pt') # Azul
create_cell_style('HeaderCellCat', '#334155', '#ffffff', bold=True, font_size='10pt') # Pizarra
create_cell_style('GuideCell', '#f8fafc', '#64748b', italic=True, font_size='8.5pt')
create_cell_style('TextCell', '#ffffff', '#0f172a', font_size='9pt', is_text_format=True)

registered_col_styles = set()
for s in auto_styles.findall(f"{{{NS['style']}}}style[@{{{NS['style']}}}family='table-column']", NS):
    registered_col_styles.add(s.attrib.get(f"{{{NS['style']}}}name"))

def get_col_style(width):
    clean_w = width.replace('.', '_').replace('cm', '')
    sname = f"col_w_{clean_w}"
    if sname not in registered_col_styles:
        c_style = ET.SubElement(auto_styles, f"{{{NS['style']}}}style", {
            f"{{{NS['style']}}}name": sname,
            f"{{{NS['style']}}}family": 'table-column'
        })
        ET.SubElement(c_style, f"{{{NS['style']}}}table-column-properties", {
            f"{{{NS['style']}}}column-width": width,
            f"{{{NS['style']}}}use-optimal-column-width": 'false'
        })
        registered_col_styles.add(sname)
    return sname

col_default_rest = get_col_style('2.5cm')

spreadsheet = root.find('.//office:spreadsheet', NS)

# Contenedor de validaciones
content_validations = spreadsheet.find(f"{{{NS['table']}}}content-validations")
if content_validations is None:
    content_validations = ET.Element(f"{{{NS['table']}}}content-validations")
    spreadsheet.insert(0, content_validations)
else:
    content_validations.clear()

def add_val_rule(name, condition_str):
    v_elem = ET.SubElement(content_validations, f"{{{NS['table']}}}content-validation", {
        f"{{{NS['table']}}}name": name,
        f"{{{NS['table']}}}condition": condition_str,
        f"{{{NS['table']}}}allow-empty-cell": 'true',
        f"{{{NS['table']}}}display-list": 'unsorted',
        f"{{{NS['table']}}}base-cell-address": 'Catalogos.A1'
    })
    ET.SubElement(v_elem, f"{{{NS['table']}}}help-message", {f"{{{NS['table']}}}display": 'false'})
    ET.SubElement(v_elem, f"{{{NS['table']}}}error-message", {f"{{{NS['table']}}}display": 'false'})

# Para listas grandes: referencia directa a la hoja Catalogos (sin límite de 255 caracteres)
add_val_rule('val_sellos', f'cell-content-is-in-list(Catalogos.$A$2:$A${len(sellos_list)+1})')
add_val_rule('val_personas', f'cell-content-is-in-list(Catalogos.$B$2:$B${len(personas_list)+1})')
add_val_rule('val_grupos', f'cell-content-is-in-list(Catalogos.$C$2:$C${len(grupos_list)+1})')

# Para listas fijas cortas: lista explícita
def make_list_cond(items):
    return 'cell-content-is-in-list(' + ';'.join([f'"{x}"' for x in items]) + ')'

add_val_rule('val_sino', make_list_cond(sino_list))
add_val_rule('val_lados', make_list_cond(lados_list))
add_val_rule('val_lados_lp', make_list_cond(lados_lp_list))
add_val_rule('val_camelot', make_list_cond(camelot_list))
add_val_rule('val_generos', make_list_cond(generos_list))

def build_ods_row(parent_table, values, style_name, validations_map=None):
    row = ET.SubElement(parent_table, f"{{{NS['table']}}}table-row")
    for col_idx, val in enumerate(values):
        attribs = {
            f"{{{NS['table']}}}style-name": style_name,
            f"{{{NS['office']}}}value-type": 'string',
            f"{{{NS['calcext']}}}value-type": 'string'
        }
        if validations_map and col_idx in validations_map:
            attribs[f"{{{NS['table']}}}content-validation-name"] = validations_map[col_idx]
            
        cell = ET.SubElement(row, f"{{{NS['table']}}}table-cell", attribs)
        p = ET.SubElement(cell, f"{{{NS['text']}}}p")
        p.text = str(val) if val is not None else ''

# Generar las 10 hojas en ODS
for sname, sdata in SHEETS_DEF.items():
    for t in list(spreadsheet.findall(f"{{{NS['table']}}}table")):
        if t.attrib.get(f"{{{NS['table']}}}name") == sname:
            spreadsheet.remove(t)
            
    tbl = ET.SubElement(spreadsheet, f"{{{NS['table']}}}table", {
        f"{{{NS['table']}}}name": sname
    })
    
    for w in sdata['widths']:
        col_sname = get_col_style(w)
        ET.SubElement(tbl, f"{{{NS['table']}}}table-column", {
            f"{{{NS['table']}}}style-name": col_sname,
            f"{{{NS['table']}}}default-cell-style-name": 'TextCell'
        })
    ET.SubElement(tbl, f"{{{NS['table']}}}table-column", {
        f"{{{NS['table']}}}style-name": col_default_rest,
        f"{{{NS['table']}}}number-columns-repeated": str(max(10, 1024 - len(sdata['widths']))),
        f"{{{NS['table']}}}default-cell-style-name": 'TextCell'
    })
    
    h_style = 'HeaderCellA' if sdata['type'] == 'ingreso' else 'HeaderCellB'
    build_ods_row(tbl, sdata['headers'], h_style)
    build_ods_row(tbl, sdata['guide'], 'GuideCell')
    
    for ex in sdata['examples']:
        build_ods_row(tbl, ex, 'TextCell', sdata['validations'])
        
    empty_rows = max(0, 50 - len(sdata['examples']))
    for _ in range(empty_rows):
        empty_vals = [''] * len(sdata['headers'])
        build_ods_row(tbl, empty_vals, 'TextCell', sdata['validations'])

# Hoja Catalogos en ODS
cat_headers = [
    'SELLOS DISCOGRÁFICOS', 'PERSONAS (COMPOSITORES)', 'GRUPOS / ORQUESTAS',
    'CLAVES CAMELOT', 'OPCIONES SI / NO', 'LADOS FÍSICOS', 'LADOS EN LP', 'GÉNEROS MUSICALES'
]
cat_widths = ['5.5cm', '6.5cm', '6.0cm', '3.5cm', '3.5cm', '3.0cm', '3.5cm', '5.0cm']

for t in list(spreadsheet.findall(f"{{{NS['table']}}}table")):
    if t.attrib.get(f"{{{NS['table']}}}name") == 'Catalogos':
        spreadsheet.remove(t)

cat_table = ET.SubElement(spreadsheet, f"{{{NS['table']}}}table", {
    f"{{{NS['table']}}}name": 'Catalogos'
})

for w in cat_widths:
    col_sname = get_col_style(w)
    ET.SubElement(cat_table, f"{{{NS['table']}}}table-column", {
        f"{{{NS['table']}}}style-name": col_sname,
        f"{{{NS['table']}}}default-cell-style-name": 'TextCell'
    })
ET.SubElement(cat_table, f"{{{NS['table']}}}table-column", {
    f"{{{NS['table']}}}style-name": col_default_rest,
    f"{{{NS['table']}}}number-columns-repeated": '1016',
    f"{{{NS['table']}}}default-cell-style-name": 'TextCell'
})

build_ods_row(cat_table, cat_headers, 'HeaderCellCat')

max_cat_rows = max(
    len(sellos_list), len(personas_list), len(grupos_list),
    len(camelot_list), len(sino_list), len(lados_list), len(lados_lp_list), len(generos_list)
)

for i in range(max_cat_rows):
    row_vals = [
        sellos_list[i] if i < len(sellos_list) else '',
        personas_list[i] if i < len(personas_list) else '',
        grupos_list[i] if i < len(grupos_list) else '',
        camelot_list[i] if i < len(camelot_list) else '',
        sino_list[i] if i < len(sino_list) else '',
        lados_list[i] if i < len(lados_list) else '',
        lados_lp_list[i] if i < len(lados_lp_list) else '',
        generos_list[i] if i < len(generos_list) else ''
    ]
    build_ods_row(cat_table, row_vals, 'TextCell')

new_content_xml = ET.tostring(root, encoding='utf-8', xml_declaration=True)

# Guardar ODS
with zipfile.ZipFile(ODS_PATH, 'w', compression=zipfile.ZIP_DEFLATED) as z_out:
    if 'mimetype' in other_files:
        z_out.writestr('mimetype', other_files['mimetype'], compress_type=zipfile.ZIP_STORED)
        del other_files['mimetype']
    z_out.writestr('content.xml', new_content_xml)
    if settings_xml:
        z_out.writestr('settings.xml', settings_xml)
    for fname, fbytes in other_files.items():
        z_out.writestr(fname, fbytes)

print('ODS generado con validaciones de rango y sin desfases.')

# ==============================================================================
# 4. CONSTRUCCIÓN DE DATOS BD.XLSX (MICROSOFT EXCEL 100% NATIVO)
# ==============================================================================
wb = openpyxl.Workbook()
# Eliminar hoja por defecto
wb.remove(wb.active)

header_fill_a = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid') # Navy
header_fill_b = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Azul
header_fill_cat = PatternFill(start_color='334155', end_color='334155', fill_type='solid') # Slate
guide_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
guide_font = Font(name='Arial', size=8.5, italic=True, color='64748B')
data_font = Font(name='Arial', size=9, color='0F172A')

border_thin = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# 1. Crear primero la hoja Catalogos en Excel para que las formulas puedan referenciarla
ws_cat = wb.create_sheet('Catalogos')
ws_cat.append(cat_headers)
for col_num in range(1, len(cat_headers) + 1):
    c = ws_cat.cell(row=1, column=col_num)
    c.fill = header_fill_cat
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for i in range(max_cat_rows):
    row_vals = [
        sellos_list[i] if i < len(sellos_list) else '',
        personas_list[i] if i < len(personas_list) else '',
        grupos_list[i] if i < len(grupos_list) else '',
        camelot_list[i] if i < len(camelot_list) else '',
        sino_list[i] if i < len(sino_list) else '',
        lados_list[i] if i < len(lados_list) else '',
        lados_lp_list[i] if i < len(lados_lp_list) else '',
        generos_list[i] if i < len(generos_list) else ''
    ]
    ws_cat.append(row_vals)
    for col_num in range(1, len(row_vals) + 1):
        c = ws_cat.cell(row=i+2, column=col_num)
        c.font = data_font
        c.number_format = '@'

# 2. Generar las 10 hojas en Excel con DataValidation nativo
for sname, sdata in SHEETS_DEF.items():
    ws = wb.create_sheet(sname)
    
    # Validaciones Excel
    dv_rules = {}
    for col_idx, vname in sdata['validations'].items():
        col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
        if vname == 'val_sellos':
            dv = DataValidation(type='list', formula1=f'=Catalogos!$A$2:$A${len(sellos_list)+1}', allow_blank=True)
        elif vname == 'val_personas':
            dv = DataValidation(type='list', formula1=f'=Catalogos!$B$2:$B${len(personas_list)+1}', allow_blank=True)
        elif vname == 'val_grupos':
            dv = DataValidation(type='list', formula1=f'=Catalogos!$C$2:$C${len(grupos_list)+1}', allow_blank=True)
        elif vname == 'val_camelot':
            dv = DataValidation(type='list', formula1='=Catalogos!$D$2:$D$25', allow_blank=True)
        elif vname == 'val_sino':
            dv = DataValidation(type='list', formula1='=Catalogos!$E$2:$E$3', allow_blank=True)
        elif vname == 'val_lados':
            dv = DataValidation(type='list', formula1='=Catalogos!$F$2:$F$3', allow_blank=True)
        elif vname == 'val_lados_lp':
            dv = DataValidation(type='list', formula1='=Catalogos!$G$2:$G$4', allow_blank=True)
        elif vname == 'val_generos':
            dv = DataValidation(type='list', formula1='=Catalogos!$H$2:$H$9', allow_blank=True)
        else:
            continue
            
        ws.add_data_validation(dv)
        dv.add(f'{col_letter}3:{col_letter}52')
        dv_rules[col_idx] = dv
        
    # Fila 1: Headers
    ws.append(sdata['headers'])
    h_fill = header_fill_a if sdata['type'] == 'ingreso' else header_fill_b
    for col_num in range(1, len(sdata['headers']) + 1):
        c = ws.cell(row=1, column=col_num)
        c.fill = h_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
    # Fila 2: Guía
    ws.append(sdata['guide'])
    for col_num in range(1, len(sdata['guide']) + 1):
        c = ws.cell(row=2, column=col_num)
        c.fill = guide_fill
        c.font = guide_font
        c.alignment = Alignment(vertical='center', wrap_text=True)
        
    # Filas 3..N: Ejemplos
    for ex in sdata['examples']:
        ws.append(ex)
        
    # Filas vacías hasta 52
    empty_rows = max(0, 50 - len(sdata['examples']))
    for _ in range(empty_rows):
        ws.append([''] * len(sdata['headers']))
        
    # Aplicar formato de texto y bordes a todas las filas de datos
    for r in range(3, 53):
        for col_num in range(1, len(sdata['headers']) + 1):
            c = ws.cell(row=r, column=col_num)
            c.font = data_font
            c.number_format = '@'
            c.border = border_thin
            
    # Ajustar anchos aproximados de columna en Excel
    for col_idx, w in enumerate(sdata['widths']):
        col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
        cm_val = float(w.replace('cm', ''))
        # 1 cm ~ 5.2 caracteres de ancho en Excel
        ws.column_dimensions[col_letter].width = max(12, int(cm_val * 5.2))

# Ajustar anchos de Catalogos
for col_idx, w in enumerate(cat_widths):
    col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
    cm_val = float(w.replace('cm', ''))
    ws_cat.column_dimensions[col_letter].width = max(15, int(cm_val * 5.2))

wb.save(XLSX_PATH)
print(f'Excel nativo creado exitosamente en: {XLSX_PATH}')
print('¡AMBOS ARCHIVOS (ODS Y XLSX) COMPLETAMENTE SINCRONIZADOS Y OPERATIVOS!')
